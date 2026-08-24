import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Color Constants
COLOR_BG_BASE = RGBColor(15, 23, 42)       # #0F172A Deep Slate
COLOR_BG_CARD = RGBColor(30, 41, 59)       # #1E293B Surface Slate
COLOR_BG_HEADER = RGBColor(11, 17, 32)     # #0B1120 Header Dark
COLOR_GOLD = RGBColor(197, 155, 39)        # #C59B27 Royal Gold
COLOR_BLUE = RGBColor(56, 189, 248)        # #38BDF8 Electric Blue
COLOR_WHITE = RGBColor(248, 250, 252)      # #F8FAFC Crisp White
COLOR_SLATE_TEXT = RGBColor(148, 163, 184) # #94A3B8 Subtitle Slate
COLOR_DARK_TEXT = RGBColor(15, 23, 42)     # #0F172A Dark text on gold
COLOR_GREEN = RGBColor(52, 211, 153)       # #34D399 Accent Green

def create_base_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG_BASE
    bg.line.fill.background()
    return slide

def add_header(slide, title_text, slide_num):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.7))
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_BG_HEADER
    header.line.color.rgb = COLOR_GOLD
    header.line.width = Pt(1)

    tf_hdr = header.text_frame
    tf_hdr.word_wrap = True
    p_hdr = tf_hdr.paragraphs[0]
    p_hdr.text = f"🏛️ GOVERNMENT MUSEUM INTERACTIVE KIOSK   |   SLIDE {slide_num} OF 15"
    p_hdr.font.size = Pt(11)
    p_hdr.font.bold = True
    p_hdr.font.color.rgb = COLOR_GOLD
    p_hdr.font.name = "Outfit"

    hdr_right = slide.shapes.add_textbox(Inches(7.5), Inches(0.1), Inches(5.6), Inches(0.5))
    tf_hr = hdr_right.text_frame
    p_hr = tf_hr.paragraphs[0]
    p_hr.alignment = PP_ALIGN.RIGHT
    p_hr.text = title_text.upper()
    p_hr.font.size = Pt(11)
    p_hr.font.bold = True
    p_hr.font.color.rgb = COLOR_BLUE
    p_hr.font.name = "Inter"

    # Persistent Bottom Bar
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.0), Inches(13.333), Inches(0.5))
    footer.fill.solid()
    footer.fill.fore_color.rgb = COLOR_BG_HEADER
    footer.line.color.rgb = COLOR_GOLD
    footer.line.width = Pt(0.5)

    tf_ft = footer.text_frame
    p_ft = tf_ft.paragraphs[0]
    p_ft.alignment = PP_ALIGN.CENTER
    p_ft.text = "Pre-Presidential Journey of Dr. Rajendra Prasad & Champaran Movement Exhibit Deck"
    p_ft.font.size = Pt(9)
    p_ft.font.color.rgb = COLOR_SLATE_TEXT
    p_ft.font.name = "Inter"

def add_card(slide, x, y, w, h, bg_color=COLOR_BG_CARD, border_color=COLOR_GOLD, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    hero_img_path = r"C:\Users\hp\.gemini\antigravity\brain\c6935483-6f7c-4a9f-b91a-5d7d7e563d75\dr_rajendra_prasad_portrait_hero_1785527107082.jpg"

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = create_base_slide(prs)
    add_header(slide1, "Title Slide & Project Overview", 1)

    # Title Card Main
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(7.5), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "EXHIBIT SPECIFICATION & PROPOSAL"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.font.name = "Inter"

    p = tf1.add_paragraph()
    p.text = "Pre-Presidential Journey of\nDr. Rajendra Prasad"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.font.name = "Outfit"

    p = tf1.add_paragraph()
    p.text = "A 55\" Interactive Touchscreen Kiosk Exhibit for Government Museum"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = COLOR_GOLD
    p.font.name = "Inter"

    if os.path.exists(hero_img_path):
        add_card(slide1, 8.6, 1.1, 4.0, 2.3, border_color=COLOR_GOLD, border_width=2)
        slide1.shapes.add_picture(hero_img_path, Inches(8.7), Inches(1.2), width=Inches(3.8), height=Inches(2.1))

    # Center CTA Banner
    cta = add_card(slide1, 0.8, 3.5, 11.733, 0.7, bg_color=COLOR_GOLD, border_color=COLOR_WHITE)
    tf_cta = cta.text_frame
    p_cta = tf_cta.paragraphs[0]
    p_cta.alignment = PP_ALIGN.CENTER
    p_cta.text = "👇 TOUCH SCREEN ANYWHERE TO BEGIN  |  अन्वेषण के लिए स्क्रीन को स्पर्श करें"
    p_cta.font.size = Pt(15)
    p_cta.font.bold = True
    p_cta.font.color.rgb = COLOR_DARK_TEXT
    p_cta.font.name = "Outfit"

    # Info Cards at Bottom
    c1 = add_card(slide1, 0.8, 4.4, 3.6, 2.3, border_color=COLOR_BLUE)
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎯 Project Target"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p = tf.add_paragraph()
    p.text = "Digital Storytelling of Dr. Rajendra Prasad's leadership in freedom struggle & nation-building prior to 1950."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_SLATE_TEXT

    c2 = add_card(slide1, 4.86, 4.4, 3.6, 2.3, border_color=COLOR_BLUE)
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🖥️ Hardware Specification"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p = tf.add_paragraph()
    p.text = "55-inch Commercial 4K Touchscreen Display with Directional Sound Dome & Fanless Industrial Mini PC."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_SLATE_TEXT

    c3 = add_card(slide1, 8.93, 4.4, 3.6, 2.3, border_color=COLOR_BLUE)
    tf = c3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🏛️ Museum Theme"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p = tf.add_paragraph()
    p.text = "Special Audio-Visual focus on Champaran Satyagraha (1917), Indigo Farmers' letters & Constituent Assembly leadership."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_SLATE_TEXT


    # ==========================================
    # SLIDE 2: Executive Summary & Museum Vision
    # ==========================================
    slide2 = create_base_slide(prs)
    add_header(slide2, "Executive Summary & Museum Vision", 2)

    # 3 Large Feature Columns
    cols_data = [
        ("🎯 Core Objective", "Immersive 9-Page Gallery", 
         ["Multi-faceted narrative before presidency (1884–1950)", 
          "Seamless touchscreen navigation for all age groups",
          "Bilingual Hindi & English instant toggle support",
          "High-definition archival document inspection"]),
        ("🌾 Key Highlight", "Champaran Satyagraha Focus", 
         ["Special Audio-Visual spotlight on 1917 movement",
          "Archiving indigo farmers' grievances & original letters",
          "Synchronized voice narration & interactive document reader",
          "First civil disobedience movement in India"]),
        ("✨ Visitor Impact", "Multi-Sensory Experience", 
         ["Transform static museum display into active exploration",
          "Educational tool for students, researchers & tourists",
          "Directional sound dome to prevent hall noise pollution",
          "60 FPS fluid touch animations & auto-reset idle timer"])
    ]

    for i, (head, subhead, bullets) in enumerate(cols_data):
        x_pos = 0.8 + i * 4.0
        c = add_card(slide2, x_pos, 1.1, 3.7, 5.6, border_color=COLOR_GOLD)
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD
        p.font.name = "Outfit"

        p_sub = tf.add_paragraph()
        p_sub.text = subhead
        p_sub.font.size = Pt(13)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_BLUE
        p_sub.font.name = "Inter"

        tf.add_paragraph() # spacing

        for bullet in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {bullet}"
            pb.font.size = Pt(11)
            pb.font.color.rgb = COLOR_WHITE
            pb.font.name = "Inter"


    # ==========================================
    # SLIDE 3: Information Architecture & Visitor Journey
    # ==========================================
    slide3 = create_base_slide(prs)
    add_header(slide3, "9-Page Information Architecture & Flow", 3)

    title_box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6))
    tf3 = title_box3.text_frame
    p = tf3.paragraphs[0]
    p.text = "SEQUENTIAL VISITOR STORYTELLING FLOW (9 INTERACTIVE PHASES)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    milestones_full = [
        ("Page 1: Welcome & Intro", "Ambient Video & Language Selection"),
        ("Page 2: The Advocate", "Calcutta & Patna High Court Practice"),
        ("Page 3: The Journalist", "Searchlight & Desh Editorial Leadership"),
        ("Page 4: The Humanitarian", "1934 Bihar Earthquake Relief Operations"),
        ("Page 5: The Educationist", "Co-founding Bihar Vidyapeeth (1921)"),
        ("Page 6: Champaran Leader", "1917 Satyagraha & Farmers' Letters (3 Parts)"),
        ("Page 7: Agriculture Minister", "1946 Interim Govt & Grow More Food"),
        ("Page 8: Constitution Maker", "Presiding President of Constituent Assembly"),
        ("Page 9: First President", "Unanimous Election & 1950 Oath Ceremony")
    ]

    for idx, (m_title, m_desc) in enumerate(milestones_full):
        row = idx // 3
        col = idx % 3
        cx = 0.8 + col * 4.0
        cy = 1.45 + row * 1.75

        card = add_card(slide3, cx, cy, 3.7, 1.55, border_color=COLOR_BLUE if idx==5 else COLOR_GOLD)
        tf_c = card.text_frame
        tf_c.word_wrap = True
        
        p1 = tf_c.paragraphs[0]
        p1.text = f"{idx+1}. {m_title}"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_BLUE if idx==5 else COLOR_WHITE
        p1.font.name = "Outfit"

        p2 = tf_c.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_SLATE_TEXT
        p2.font.name = "Inter"

        p3 = tf_c.add_paragraph()
        p3.text = "➜ Tap to Explore Node"
        p3.font.size = Pt(9)
        p3.font.italic = True
        p3.font.color.rgb = COLOR_GOLD


    # ==========================================
    # SLIDE 4: Pages 1 & 2 – Welcome & "The Advocate"
    # ==========================================
    slide4 = create_base_slide(prs)
    add_header(slide4, "Page 1 & 2: Introduction & Legal Luminary", 4)

    # Left Card Page 1
    c_p1 = add_card(slide4, 0.8, 1.1, 5.7, 5.6, border_color=COLOR_GOLD)
    tf1 = c_p1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🖥️ Page 1: Welcome & Landing Screen"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    p_items = [
        "Ambient Visuals: High-definition archival loop of Bihar freedom struggle",
        "Language Selection: Instant Hindi / English UI & Voice toggle buttons",
        "Attract Mode Prompt: Glowing 'Touch Screen Anywhere to Begin' animation",
        "Kiosk Safety: Automatic idle timer resets to Page 1 after 60s inactivity"
    ]
    for item in p_items:
        p = tf1.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE

    # Right Card Page 2
    c_p2 = add_card(slide4, 6.8, 1.1, 5.7, 5.6, border_color=COLOR_BLUE)
    tf2 = c_p2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "⚖️ Page 2: The Advocate (Legal Career)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p_items2 = [
        "High Court Practice: Legal luminary at Calcutta High Court (1911) & Patna High Court (1916)",
        "Defending Freedom Fighters: Provided pro-bono legal counsel to activists and civil rights causes",
        "Interactive Feature: Tap-to-expand major court cases & legal arguments",
        "Document Zoom Viewer: Inspect high-res digital scans of original law certificates & legal briefs"
    ]
    for item in p_items2:
        p = tf2.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE


    # ==========================================
    # SLIDE 5: Pages 3 & 4 – "The Journalist" & "The Humanitarian"
    # ==========================================
    slide5 = create_base_slide(prs)
    add_header(slide5, "Page 3 & 4: Press Freedom & Disaster Relief", 5)

    c_p3 = add_card(slide5, 0.8, 1.1, 5.7, 5.6, border_color=COLOR_GOLD)
    tf3 = c_p3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "📰 Page 3: The Journalist (Press Freedom)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    p_items3 = [
        "Editorial Leadership: Editor & contributor to 'Searchlight' (English) & 'Desh' (Hindi) newspapers",
        "Spreading Consciousness: Used journalism as a non-violent instrument for nationalist awakening",
        "Interactive Flip-Book: Touch gesture digital newspaper reader showing digitized 1920s front pages",
        "Archival Overlays: Pop-up translations of key patriotic editorials written by Dr. Prasad"
    ]
    for item in p_items3:
        p = tf3.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE

    c_p4 = add_card(slide5, 6.8, 1.1, 5.7, 5.6, border_color=COLOR_BLUE)
    tf4 = c_p4.text_frame
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    p.text = "🤝 Page 4: The Humanitarian (1934 Relief)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p_items4 = [
        "1934 Bihar Earthquake: Led massive relief operations following devastating 8.0 magnitude quake",
        "Bihar Central Relief Committee: Raised over 3.8 million rupees while imprisoned & released to lead",
        "Interactive Relief Map: Touch districts (Monghyr, Muzaffarpur, Patna) to view archival damage photos",
        "Historical Photography: High-res gallery of community kitchens, shelter construction & medical aid"
    ]
    for item in p_items4:
        p = tf4.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE


    # ==========================================
    # SLIDE 6: Page 5 – "The Educationist"
    # ==========================================
    slide6 = create_base_slide(prs)
    add_header(slide6, "Page 5: Shaping National Education (Bihar Vidyapeeth)", 6)

    c_p5_left = add_card(slide6, 0.8, 1.1, 7.0, 5.6, border_color=COLOR_GOLD)
    tf5 = c_p5_left.text_frame
    tf5.word_wrap = True
    p = tf5.paragraphs[0]
    p.text = "🎓 Page 5: Indigenous Education & Bihar Vidyapeeth"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    p_items5 = [
        "Swadeshi Education Vision: Promoting indigenous, self-reliant education during 1921 Non-Cooperation Movement",
        "Founding Bihar Vidyapeeth: Established in Patna (1921) inaugurated by Mahatma Gandhi & Dr. Rajendra Prasad",
        "Principal & Vice-Chancellor: Shaped curriculum combining academics, vocational skills & nationalist values",
        "Interactive Timeline: Flowchart of national educational institutions established across India in the 1920s",
        "Audio Quote Player: Listen to Dr. Prasad's speeches on youth empowerment and character building"
    ]
    for item in p_items5:
        p = tf5.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE

    c_p5_right = add_card(slide6, 8.1, 1.1, 4.4, 5.6, border_color=COLOR_BLUE)
    tf5_r = c_p5_right.text_frame
    tf5_r.word_wrap = True
    p = tf5_r.paragraphs[0]
    p.text = "🖼️ Archival Feature Display"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p = tf5_r.add_paragraph()
    p.text = "Inauguration Photo & Audio Widget"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_GOLD

    p = tf5_r.add_paragraph()
    p.text = "• Interactive Photo Gallery: Rare photographs of Bihar Vidyapeeth campus, faculty & early batches.\n\n• Audio Narration: Directional audio overlay providing historical commentary in Hindi & English."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_SLATE_TEXT


    # ==========================================
    # SLIDE 7: Page 6 (Part 1) – "The Champaran Movement & Gandhiji"
    # ==========================================
    slide7 = create_base_slide(prs)
    add_header(slide7, "Page 6 (Part 1): Role in Champaran Satyagraha (1917)", 7)

    c_p6_1 = add_card(slide7, 0.8, 1.1, 11.733, 5.6, border_color=COLOR_GOLD)
    tf6_1 = c_p6_1.text_frame
    tf6_1.word_wrap = True

    p = tf6_1.paragraphs[0]
    p.text = "🌱 The Turning Point: Champaran Satyagraha (1917)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    p_items6_1 = [
        "First Association with Gandhiji: Dr. Rajendra Prasad met Mahatma Gandhi in April 1917 in Patna & Champaran, beginning a lifelong partnership.",
        "Logistics & Legal Operations: Managed the massive task of organizing volunteer teams, legal defense, and accommodation across North Bihar villages.",
        "Language & Dialect Translation: Translated local Bhojpuri and Maithili testimonies into English for British inquiry committees and legal records.",
        "Recording Statement Archives: Personally recorded and verified thousands of tenant farmer statements detailing colonial oppression.",
        "National Impact: Laying the foundation for India's first successful Civil Disobedience campaign, leading to the Champaran Agrarian Act of 1918."
    ]

    for item in p_items6_1:
        p = tf6_1.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_WHITE
        tf6_1.add_paragraph() # spacing


    # ==========================================
    # SLIDE 8: Page 6 (Part 2) – Deep-Dive: Indigo Farmers' Struggles
    # ==========================================
    slide8 = create_base_slide(prs)
    add_header(slide8, "Feature Spotlight: Indigo Farmers' Grievances", 8)

    grid_data = [
        ("⛓️ The Tinkathia System", "Colonial Exploitation", "Forced cultivation of indigo on 3/20th (3 kathas per bigha) of every tenant farmer's best fertile land for European planters."),
        ("💸 Excessive Taxation & Extortion", "Economic Ruin", "Illegal cesses (Abwabs), arbitrary fines, and fixed low prices paid for indigo, forcing peasant families into perpetual debt."),
        ("🌾 Land & Soil Degradation", "Famine Risk", "Indigo crops depleted soil nutrients rapidly, preventing farmers from growing food grains like rice and wheat essential for survival."),
        ("📜 Archival Evidence Collection", "7,000+ Recorded Letters", "Dr. Rajendra Prasad & team systematically cataloged thousands of signed farmer grievances that exposed planter atrocities.")
    ]

    for i, (title, subtitle, desc) in enumerate(grid_data):
        row = i // 2
        col = i % 2
        cx = 0.8 + col * 6.0
        cy = 1.1 + row * 2.85

        card = add_card(slide8, cx, cy, 5.7, 2.65, border_color=COLOR_BLUE if i%2==1 else COLOR_GOLD)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD if i%2==0 else COLOR_BLUE

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_SLATE_TEXT

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_WHITE


    # ==========================================
    # SLIDE 9: Page 6 (Part 3) – AV & Interactive Letter Explorer
    # ==========================================
    slide9 = create_base_slide(prs)
    add_header(slide9, "Feature Spotlight: Interactive Letter & Audio Sync", 9)

    c_av1 = add_card(slide9, 0.8, 1.1, 3.7, 5.6, border_color=COLOR_GOLD)
    tf = c_av1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎙️ Professional Voice Narration"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    p_b = tf.add_paragraph()
    p_b.text = "• Dramatic readings of translated farmer letters by professional voice actors.\n\n• Subtle ambient rural soundscapes (birds, charkha, rain) for deep immersion.\n\n• Directional audio overhead dome control."
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = COLOR_WHITE

    c_av2 = add_card(slide9, 4.8, 1.1, 3.7, 5.6, border_color=COLOR_BLUE)
    tf = c_av2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⏱️ Synchronized Karaoke Highlight"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p_b = tf.add_paragraph()
    p_b.text = "• Real-time text highlighting on screen synchronized line-by-line with audio narration.\n\n• Enhanced accessibility for hearing impaired visitors and young students.\n\n• Interactive speed control."
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = COLOR_WHITE

    c_av3 = add_card(slide9, 8.8, 1.1, 3.7, 5.6, border_color=COLOR_GREEN)
    tf = c_av3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📂 Interactive Letter Browser"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    p_b = tf.add_paragraph()
    p_b.text = "• Categorized filter tabs:\n  - Taxation & Fines\n  - Physical Coercion\n  - Personal Testimonies\n\n• Touch any thumbnail to view high-res handwritten manuscript scan with Hindi/English side-by-side translation."
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = COLOR_WHITE


    # ==========================================
    # SLIDE 10: Page 7 – "The Agriculture Minister"
    # ==========================================
    slide10 = create_base_slide(prs)
    add_header(slide10, "Page 7: Food Security & Agriculture (1946)", 10)

    c_ag_l = add_card(slide10, 0.8, 1.1, 7.0, 5.6, border_color=COLOR_GOLD)
    tf = c_ag_l.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🌾 Minister for Food & Agriculture (1946 Interim Govt)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    ag_items = [
        "Post-War Food Crisis: Appointed Minister in Interim Government (September 1946) during severe post-WWII food shortages and Bengal famine aftermath.",
        "'Grow More Food' Campaign: Spearheaded nationwide agricultural campaign to expand cultivated area, improve seed distribution and irrigation works.",
        "Dynamic Data Infographics: Interactive Touch UI charts depicting grain distribution, food imports reduction and regional production statistics.",
        "Newsreel Archive: Touch player featuring historical 1946 newsreel footage of agricultural rallies and storage grain silo inaugurations."
    ]
    for item in ag_items:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE
        tf.add_paragraph()

    c_ag_r = add_card(slide10, 8.1, 1.1, 4.4, 5.6, border_color=COLOR_BLUE)
    tf_r = c_ag_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "📊 Interactive Kiosk Widget"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p = tf_r.add_paragraph()
    p.text = "Yield Statistics Slider"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_GOLD

    p = tf_r.add_paragraph()
    p.text = "• Interactive Year Slider (1946–1950):\nAllows visitors to drag across timeline and see agricultural yield data, land reform bills, and food grain self-reliance milestones."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_SLATE_TEXT


    # ==========================================
    # SLIDE 11: Page 8 – "The Constitution Maker"
    # ==========================================
    slide11 = create_base_slide(prs)
    add_header(slide11, "Page 8: President of the Constituent Assembly", 11)

    c_const = add_card(slide11, 0.8, 1.1, 11.733, 5.6, border_color=COLOR_GOLD)
    tf = c_const.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📜 Presiding over the Birth of the Republic (1946–1950)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    const_items = [
        "Unanimous Election: Elected President of the Constituent Assembly of India on December 11, 1946.",
        "Guiding Assembly Debates: Impartially steered nearly 3 years of intense debates among 299 assembly members to forge the world's longest written constitution.",
        "Harmonizing Diverse Views: Known for his patience, dignity, and legal acumen in resolving contentious issues on federalism, fundamental rights, and official languages.",
        "High-Res Constitution Viewer: Ultra-HD digital scan of the original illustrated Constitution of India (hand-written by Prem Behari Narain Raizada and decorated by Nandalal Bose).",
        "Interactive Member Signature Map: Touch assembly members' signatures to view their biographies, speeches, and state representation."
    ]
    for item in const_items:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE
        tf.add_paragraph()


    # ==========================================
    # SLIDE 12: Page 9 – Culmination: First President of India (1950)
    # ==========================================
    slide12 = create_base_slide(prs)
    add_header(slide12, "Page 9: Election as First President (1950)", 12)

    c_pres_l = add_card(slide12, 0.8, 1.1, 6.8, 5.6, border_color=COLOR_GOLD)
    tf = c_pres_l.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🇮🇳 Dawn of the Sovereign Republic"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    pres_items = [
        "Unanimous Election: Unanimously elected as the first President of India on January 24, 1950, taking oath on January 26, 1950 (Republic Day).",
        "Setting Democratic Standards: Established enduring presidential traditions of humility, non-partisan dignity, accessibility, and high ethical standards.",
        "Longest Serving President: Remained in office for 12 years (1950–1962), guiding the nascent democracy through three general elections.",
        "Interactive Oath Document: High-res digital replica of the original 1950 Swearing-in Declaration with tap-to-translate text."
    ]
    for item in pres_items:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE
        tf.add_paragraph()

    c_pres_r = add_card(slide12, 7.9, 1.1, 4.6, 5.6, border_color=COLOR_BLUE)
    tf_r = c_pres_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "🎥 Archival Video Archive"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p = tf_r.add_paragraph()
    p.text = "1950 Republic Day Oath Ceremony"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_GOLD

    p = tf_r.add_paragraph()
    p.text = "• Interactive Video Player:\nFeatures restored black-and-white film footage of the swearing-in ceremony at Durbar Hall, Government House, and 31-gun salute."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_SLATE_TEXT


    # ==========================================
    # SLIDE 13: UI/UX & Interactive Design System
    # ==========================================
    slide13 = create_base_slide(prs)
    add_header(slide13, "Touch-First UI/UX & Accessibility Standards", 13)

    ux_cards = [
        ("👆 Touch Ergonomics", "55\" Kiosk Layout", "Optimized target touch zones (minimum 80x80px) positioned between 3.5ft and 5.5ft height for comfortable reach by adults and children."),
        ("🌐 Bilingual Interface", "Instant Toggle", "One-tap switching between Hindi and English for all UI labels, document translations, and voice narration audio tracks."),
        ("⏱️ Inactivity Auto-Reset", "Idle Management", "Automatic 60-second inactivity timer resets kiosk to Page 1 Welcome Screen to welcome new museum visitors."),
        ("🔒 Kiosk Lockdown & 60 FPS", "System Security", "Kiosk mode restricts OS access; hardware-accelerated CSS/JS animations guarantee silky 60 FPS touch responsiveness.")
    ]

    for i, (title, subtitle, desc) in enumerate(ux_cards):
        row = i // 2
        col = i % 2
        cx = 0.8 + col * 6.0
        cy = 1.1 + row * 2.85

        card = add_card(slide13, cx, cy, 5.7, 2.65, border_color=COLOR_BLUE if i%2==1 else COLOR_GOLD)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD if i%2==0 else COLOR_BLUE

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_SLATE_TEXT

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_WHITE


    # ==========================================
    # SLIDE 14: Hardware & Audio Integration Specs
    # ==========================================
    slide14 = create_base_slide(prs)
    add_header(slide14, "Hardware Specifications & Audio Setup", 14)

    hw_specs = [
        ("🖥️ Commercial 4K Touch Display", "55\" Ultra-HD 3840x2160 IPS Panel, 500 nits brightness, 10-point Infrared/Capacitive touch, anti-glare toughened glass."),
        ("⚙️ Industrial Mini PC Engine", "Intel Core i5/i7 Industrial PC, 16GB DDR4 RAM, 512GB NVMe SSD, Fanless enclosure for 24/7 continuous museum operation."),
        ("🔊 Directional Audio Sound Dome", "Overhead directional speaker dome / focused sound beam that provides crystal clear audio to kiosk user without bleed into gallery."),
        ("🏗️ Custom Museum Enclosure", "Heavy-duty powder-coated steel kiosk housing with concealed cable management, key-locked maintenance access & ventilation.")
    ]

    for i, (title, desc) in enumerate(hw_specs):
        cy = 1.1 + i * 1.4
        card = add_card(slide14, 0.8, cy, 11.733, 1.25, border_color=COLOR_GOLD if i%2==0 else COLOR_BLUE)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD if i%2==0 else COLOR_BLUE

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_WHITE


    # ==========================================
    # SLIDE 15: Implementation Roadmap & Deliverables
    # ==========================================
    slide15 = create_base_slide(prs)
    add_header(slide15, "Project Implementation Roadmap & Deliverables", 15)

    phases = [
        ("Phase 1: Content & Audio (W1-W2)", "Archival Curation & Scripting", ["Historical document & photo curation", "Hindi/English voiceover recording", "Letter transcript translation"]),
        ("Phase 2: UI/UX & Software (W3-W4)", "Software Development", ["55\" Touch UI/UX design tokens", "Interactive HTML5/JS app engine", "Audio-text sync module"]),
        ("Phase 3: Hardware & QA (W5)", "Integration & Testing", ["Mini PC & Touchscreen assembly", "Directional audio calibration", "60 FPS stress testing & QA"]),
        ("Phase 4: On-Site Launch (W6)", "Installation & Training", ["Museum kiosk site installation", "Staff maintenance training", "Final handover & warranty launch"])
    ]

    for i, (p_title, p_sub, p_bullets) in enumerate(phases):
        cx = 0.8 + i * 3.0
        card = add_card(slide15, cx, 1.1, 2.7, 5.6, border_color=COLOR_GOLD if i%2==0 else COLOR_BLUE)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD if i%2==0 else COLOR_BLUE

        p_subhead = tf.add_paragraph()
        p_subhead.text = p_sub
        p_subhead.font.size = Pt(11)
        p_subhead.font.bold = True
        p_subhead.font.color.rgb = COLOR_SLATE_TEXT

        tf.add_paragraph()

        for b in p_bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.size = Pt(10)
            pb.font.color.rgb = COLOR_WHITE

    # Save presentation to multiple convenient locations
    out1 = r"d:\Museum project\Dr_Rajendra_Prasad_15_Slide_Interactive_Kiosk_Exhibit.pptx"
    out2 = r"d:\Museum project\Slide1_Dr_Rajendra_Prasad_Kiosk.pptx"
    downloads_path = r"C:\Users\hp\Downloads\Dr_Rajendra_Prasad_15_Slide_Interactive_Kiosk_Exhibit.pptx"

    prs.save(out1)
    prs.save(out2)
    
    # Save to Downloads if possible
    try:
        downloads_dir = r"C:\Users\hp\Downloads"
        if os.path.exists(downloads_dir):
            prs.save(downloads_path)
            print(f"Saved to Downloads: {downloads_path}")
    except Exception as e:
        print(f"Could not save to Downloads: {e}")

    print(f"Saved successfully to: {out1} and {out2}")

if __name__ == "__main__":
    build_presentation()
